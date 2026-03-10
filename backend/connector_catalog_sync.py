import re
import uuid
from pathlib import Path

import models
from database import SessionLocal


WORKBOOK_NAME = "Pre-built Connectors - New Use Cases Added Feb 26 V1.xlsx"


def _slugify_connector_name(name: str) -> str:
    normalized = re.sub(r"[^a-zA-Z0-9]+", "-", name.strip().lower()).strip("-")
    return normalized or f"connector-{uuid.uuid4().hex[:8]}"


def _resolve_workbook_path(workbook_name: str = WORKBOOK_NAME) -> Path | None:
    candidates = [
        Path(__file__).resolve().parents[1] / "docs" / workbook_name,
        Path(__file__).resolve().parents[2] / workbook_name,
        Path(__file__).resolve().parents[1] / workbook_name,
        Path.cwd() / workbook_name,
    ]
    return next((path for path in candidates if path.exists()), None)


def load_connector_catalog_from_excel(workbook_path: Path | None = None) -> list[dict]:
    resolved_path = workbook_path or _resolve_workbook_path()
    if not resolved_path:
        print(f"[catalog] workbook not found: {WORKBOOK_NAME}")
        return []

    try:
        import openpyxl  # type: ignore
    except Exception as exc:
        print(f"[catalog] openpyxl unavailable: {exc}")
        return []

    wb = openpyxl.load_workbook(resolved_path, data_only=True, read_only=True)
    ws = wb["Template"] if "Template" in wb.sheetnames else wb[wb.sheetnames[0]]

    # Read header row to find use case column(s)
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    headers = list(header_row) if header_row else []
    usecase_col_indices: list[int] = []
    for i, h in enumerate(headers):
        if h is None:
            continue
        h_lower = str(h).lower().strip()
        if "use" in h_lower and "case" in h_lower:
            usecase_col_indices.append(i)
        elif h_lower in ("ingestion", "action", "usecase"):
            usecase_col_indices.append(i)
    usecase_from_header = bool(usecase_col_indices)
    if not usecase_col_indices:
        usecase_col_indices = [2, 3, 4, 5]

    rows: list[dict] = []
    used_ids: set[str] = set()
    current_type = ""
    for row_tuple in ws.iter_rows(min_row=2, values_only=True):
        parts = list(row_tuple) if row_tuple else []
        category = parts[0] if len(parts) > 0 else None
        system_name = parts[1] if len(parts) > 1 else None
        usecase_parts: list[str] = []
        for idx in usecase_col_indices:
            if idx < len(parts) and parts[idx] is not None:
                val = str(parts[idx]).strip()
                if val:
                    usecase_parts.append(val)
                    if not usecase_from_header:
                        break
        usecase_val = "\n".join(usecase_parts) if usecase_parts else None
        if category and str(category).strip():
            current_type = str(category).strip()

        if not system_name or not str(system_name).strip():
            continue

        name = str(system_name).strip()
        if name.lower() == "system name":
            continue

        base_id = _slugify_connector_name(name)
        connector_id = base_id
        suffix = 2
        while connector_id in used_ids:
            connector_id = f"{base_id}-{suffix}"
            suffix += 1
        used_ids.add(connector_id)

        rows.append(
            {
                "connector_id": connector_id,
                "name": name,
                "type": current_type or "Unknown",
                "usecase": usecase_val,
                "guide_url": f"/integration/connectors/{connector_id}/guide",
                "json_url": f"/integration/connectors/{connector_id}/json",
                "version_name": "v1.0.0",
                "logo_url": "https://via.placeholder.com/28",
            }
        )

    return rows


def sync_connector_catalog_from_excel(workbook_path: Path | None = None) -> int:
    db = SessionLocal()
    try:
        rows = load_connector_catalog_from_excel(workbook_path=workbook_path)
        if not rows:
            print("[catalog] no rows loaded from excel")
            return 0

        existing_rows = db.query(models.ConnectorCatalog).all()
        existing_map = {
            row.connector_id: (row.name, row.type, row.usecase)
            for row in existing_rows
        }
        incoming_map = {
            row["connector_id"]: (row["name"], row["type"], row.get("usecase"))
            for row in rows
        }
        # Backfill guide_url, json_url for existing rows if missing
        needs_backfill = any(
            r.guide_url is None or r.guide_url in ("Guide", "")
            or r.json_url is None or r.json_url in ("JSON", "")
            for r in existing_rows
        )
        if needs_backfill and existing_map == incoming_map:
            for r in existing_rows:
                if r.guide_url is None or r.guide_url in ("Guide", ""):
                    r.guide_url = f"/integration/connectors/{r.connector_id}/guide"
                if r.json_url is None or r.json_url in ("JSON", ""):
                    r.json_url = f"/integration/connectors/{r.connector_id}/json"
                if r.version_name is None:
                    r.version_name = "v1.0.0"
                if r.logo_url is None:
                    r.logo_url = "https://via.placeholder.com/28"
            db.commit()
            print("[catalog] backfilled guide_url, json_url for existing rows")
            return len(rows)
        if existing_map == incoming_map:
            print(f"[catalog] unchanged ({len(rows)} rows)")
            return len(rows)

        db.query(models.ConnectorCatalog).delete()
        for row in rows:
            db.add(
                models.ConnectorCatalog(
                    id=str(uuid.uuid4()),
                    connector_id=row["connector_id"],
                    name=row["name"],
                    type=row["type"],
                    usecase=row.get("usecase"),
                    guide_url=row.get("guide_url", f"/integration/connectors/{row['connector_id']}/guide"),
                    json_url=row.get("json_url", f"/integration/connectors/{row['connector_id']}/json"),
                    version_name=row.get("version_name", "v1.0.0"),
                    logo_url=row.get("logo_url", "https://via.placeholder.com/28"),
                )
            )
        db.commit()
        action = "seeded" if not existing_rows else "synchronized"
        print(f"[catalog] {action} {len(rows)} connector catalog rows")
        return len(rows)
    except Exception as exc:
        db.rollback()
        print(f"[catalog] failed seeding connector catalog: {exc}")
        raise
    finally:
        db.close()
