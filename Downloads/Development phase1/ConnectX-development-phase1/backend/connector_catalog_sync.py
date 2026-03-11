import re
import uuid
from pathlib import Path
import os

import models
from database import SessionLocal


# Single canonical workbook path for connector catalog imports.
# Put the Excel file here and run: `python connector_catalog_sync.py`
DEFAULT_WORKBOOK_PATH = Path(__file__).resolve().parent / "docs" / "connector_catalog.xlsx"
LEGACY_WORKBOOK_NAME = "Pre-built Connectors - New Use Cases Added Feb 26 V1.xlsx"


def _slugify_connector_name(name: str) -> str:
    normalized = re.sub(r"[^a-zA-Z0-9]+", "-", name.strip().lower()).strip("-")
    return normalized or f"connector-{uuid.uuid4().hex[:8]}"


def _normalize_name(value: str | None) -> str:
    return str(value or "").strip().lower()


def _resolve_workbook_path(workbook_path: Path | None = None) -> Path | None:
    if workbook_path and workbook_path.exists():
        return workbook_path

    env_path = os.getenv("CONNECTOR_CATALOG_EXCEL_PATH")
    if env_path:
        env_resolved = Path(env_path).expanduser()
        if env_resolved.exists():
            return env_resolved

    if DEFAULT_WORKBOOK_PATH.exists():
        return DEFAULT_WORKBOOK_PATH

    legacy_path = Path(__file__).resolve().parent / "docs" / LEGACY_WORKBOOK_NAME
    if legacy_path.exists():
        return legacy_path

    return None


def _header_to_field(header: str) -> str | None:
    """Map Excel header to our schema field name."""
    h = str(header or "").lower().strip()
    if not h:
        return None
    # Name / system name
    if h in ("system name", "name", "connector name", "product name"):
        return "name"
    if h in ("category", "type"):
        return "type"
    if h in ("connector_id", "connector id", "id"):
        return "connector_id"
    if "use" in h and "case" in h or h in ("ingestion", "action", "usecase"):
        return "usecase"
    return None


def load_connector_catalog_from_excel(workbook_path: Path | None = None) -> list[dict]:
    resolved_path = _resolve_workbook_path(workbook_path=workbook_path)
    if not resolved_path:
        print(
            "[catalog] workbook not found. Place Excel at "
            f"{DEFAULT_WORKBOOK_PATH} or set CONNECTOR_CATALOG_EXCEL_PATH."
        )
        return []

    try:
        import openpyxl  # type: ignore
    except Exception as exc:
        print(f"[catalog] openpyxl unavailable: {exc}")
        return []

    wb = openpyxl.load_workbook(resolved_path, data_only=True, read_only=True)
    ws = wb["Template"] if "Template" in wb.sheetnames else wb[wb.sheetnames[0]]

    # Build column index -> field mapping from header row
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    headers = list(header_row) if header_row else []
    col_map: dict[int, str] = {}  # index -> field name
    usecase_indices: list[int] = []
    for i, h in enumerate(headers):
        field = _header_to_field(h)
        if field:
            if field == "usecase":
                usecase_indices.append(i)
            else:
                col_map[i] = field

    # Fallback: if no header mapping, use positional (col 0=type, 1=name)
    if not col_map and not usecase_indices and len(headers) >= 2:
        col_map = {0: "type", 1: "name"}
        usecase_indices = [2, 3, 4, 5] if len(headers) > 5 else list(range(2, len(headers)))

    rows: list[dict] = []
    used_ids: set[str] = set()
    current_type = ""

    for row_tuple in ws.iter_rows(min_row=2, values_only=True):
        parts = list(row_tuple) if row_tuple else []
        row_data: dict[str, str | None] = {}

        # Read mapped columns
        for idx, field in col_map.items():
            if idx < len(parts) and parts[idx] is not None:
                val = str(parts[idx]).strip()
                if val:
                    row_data[field] = val

        # Track type from category column (for positional fallback)
        if "type" in row_data and row_data["type"]:
            current_type = row_data["type"]

        # Collect use case columns (can have multiple: Ingestion, Action, etc.)
        usecase_parts: list[str] = []
        for idx in usecase_indices:
            if idx < len(parts) and parts[idx] is not None:
                val = str(parts[idx]).strip()
                if val:
                    usecase_parts.append(val)
        if usecase_parts:
            row_data["usecase"] = "\n".join(usecase_parts)

        # Name is required
        name = (row_data.get("name") or "").strip()
        if not name or name.lower() in ("system name", "name"):
            continue

        # Connector ID: use from Excel or generate from name
        connector_id = (row_data.get("connector_id") or "").strip()
        if not connector_id:
            base_id = _slugify_connector_name(name)
            connector_id = base_id
            suffix = 2
            while connector_id in used_ids:
                connector_id = f"{base_id}-{suffix}"
                suffix += 1
        used_ids.add(connector_id)

        rows.append(
            {
                "connector_id": connector_id,
                "name": name,
                "type": (row_data.get("type") or current_type or "Unknown").strip(),
                "usecase": row_data.get("usecase"),
                # Not in Excel - always generated
                "guide_url": f"/integration/connectors/{connector_id}/guide",
                "json_url": f"/integration/connectors/{connector_id}/json",
                "version_name": "v1.0.0",
                "logo_url": "https://via.placeholder.com/28",
            }
        )

    return rows


def sync_connector_catalog_from_excel(workbook_path: Path | None = None) -> int:
    db = SessionLocal()
    try:
        rows = load_connector_catalog_from_excel(workbook_path=workbook_path)
        if not rows:
            print("[catalog] no rows loaded from excel")
            return 0

        existing_rows = db.query(models.ConnectorCatalog).all()
        existing_by_name: dict[str, models.ConnectorCatalog] = {}
        used_connector_ids = {row.connector_id for row in existing_rows if row.connector_id}
        for row in existing_rows:
            key = _normalize_name(row.name)
            if key and key not in existing_by_name:
                existing_by_name[key] = row

        # Deduplicate incoming rows by connector name (last one wins).
        incoming_by_name: dict[str, dict] = {}
        for row in rows:
            key = _normalize_name(row.get("name"))
            if key:
                incoming_by_name[key] = row

        created = 0
        updated = 0

        for name_key, row in incoming_by_name.items():
            existing = existing_by_name.get(name_key)
            if existing:
                changed = False
                if existing.name != row["name"]:
                    existing.name = row["name"]
                    changed = True
                if existing.type != row["type"]:
                    existing.type = row["type"]
                    changed = True
                if existing.usecase != row.get("usecase"):
                    existing.usecase = row.get("usecase")
                    changed = True

                desired_guide = f"/integration/connectors/{existing.connector_id}/guide"
                desired_json = f"/integration/connectors/{existing.connector_id}/json"
                desired_version = row.get("version_name") or "v1.0.0"
                desired_logo = row.get("logo_url") or "https://via.placeholder.com/28"

                if existing.guide_url != desired_guide:
                    existing.guide_url = desired_guide
                    changed = True
                if existing.json_url != desired_json:
                    existing.json_url = desired_json
                    changed = True
                if existing.version_name != desired_version:
                    existing.version_name = desired_version
                    changed = True
                if existing.logo_url != desired_logo:
                    existing.logo_url = desired_logo
                    changed = True

                if changed:
                    updated += 1
                continue

            requested_id = str(row.get("connector_id") or "").strip()
            connector_id = requested_id or _slugify_connector_name(row["name"])
            base_id = connector_id
            suffix = 2
            while connector_id in used_connector_ids:
                connector_id = f"{base_id}-{suffix}"
                suffix += 1
            used_connector_ids.add(connector_id)

            db.add(
                models.ConnectorCatalog(
                    id=str(uuid.uuid4()),
                    connector_id=connector_id,
                    name=row["name"],
                    type=row["type"],
                    usecase=row.get("usecase"),
                    guide_url=f"/integration/connectors/{connector_id}/guide",
                    json_url=f"/integration/connectors/{connector_id}/json",
                    version_name=row.get("version_name", "v1.0.0"),
                    logo_url=row.get("logo_url", "https://via.placeholder.com/28"),
                )
            )
            created += 1

        db.commit()
        print(
            "[catalog] upserted by name: "
            f"{len(incoming_by_name)} processed, {created} created, {updated} updated"
        )
        return len(incoming_by_name)
    except Exception as exc:
        db.rollback()
        print(f"[catalog] failed seeding connector catalog: {exc}")
        raise
    finally:
        db.close()


if __name__ == "__main__":
    synced_rows = sync_connector_catalog_from_excel()
    print(f"[catalog] done: {synced_rows} rows")
